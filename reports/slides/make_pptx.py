"""Build OR_Final_Presentation.pptx from the rendered slide PDF: one full-bleed 1920x1080 image per slide,
with the speaker notes taken from the deck's slide files.

usage: python make_pptx.py <deck_project_dir>   (the same folder export_slides.py uses)
The text on the slides is not editable in this file; an editable version can be exported from the
online deck (Share > Export > PowerPoint).
"""
import io
import json
import re
import sys
from html import unescape
from pathlib import Path

import fitz                      # PyMuPDF
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECK = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'deck' / 'project'
PDF = HERE / 'OR_Final_Presentation.pdf'
OUT = HERE / 'OR_Final_Presentation.pptx'

order = json.loads((DECK / 'deck.json').read_text(encoding='utf-8'))['order']
notes = []
for sid in order:
    h = (DECK / 'slides' / f'{sid}.html').read_text(encoding='utf-8')
    m = re.search(r'<aside>(.*?)</aside>', h, flags=re.S)
    notes.append(unescape(re.sub(r'<[^>]+>', '', m.group(1))).strip() if m else '')

doc = fitz.open(PDF)
assert len(doc) == len(order), (len(doc), len(order))
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)      # 13.333 x 7.5 in (16:9)
blank = prs.slide_layouts[6]
for page, note in zip(doc, notes):
    zoom = 1920 / page.rect.width
    png = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom)).tobytes('png')
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(io.BytesIO(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    s.notes_slide.notes_text_frame.text = note
cp = prs.core_properties
cp.title = 'Right Place, Right Time: EV Charging in Shenzhen'
cp.author = cp.last_modified_by = 'OR 76000093 final project team'
prs.save(OUT)
print('ok', len(order), 'slides ->', OUT.name)
